from __future__ import annotations
import io
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils.dataframe import dataframe_to_rows

_RED = "FFEF4444"
_ORANGE = "FFF97316"
_GREEN = "FF22C55E"
_INDIGO = "FF6366F1"
_YELLOW = "FFEAB308"
_GRAY = "FF6B7280"
_HEADER_BG = "FF1E1B4B"
_HEADER_FG = "FFFFFFFF"


def _style_header(ws, row_num: int = 1, fill_hex: str = _HEADER_BG) -> None:
    fill = PatternFill(fill_type="solid", fgColor=fill_hex)
    font = Font(bold=True, color=_HEADER_FG)
    for cell in ws[row_num]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center")


def _auto_width(ws) -> None:
    for col in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 50)


def _df_to_sheet(ws, df: pd.DataFrame, title: str, fill_hex: str = _HEADER_BG) -> None:
    ws.title = title
    for r in dataframe_to_rows(df, index=False, header=True):
        ws.append(r)
    _style_header(ws, fill_hex=fill_hex)
    _auto_width(ws)


def build(ppc: dict, seo: dict) -> bytes:
    wb = Workbook()

    # --- Executive Summary ---
    ws_exec = wb.active
    ws_exec.title = "Executive Summary"
    rows = [
        ["Komacut Growth Audit — Executive Summary", ""],
        ["", ""],
        ["PPC FINDINGS", ""],
        ["Wasted Budget (zero-conversion keywords)", f"${ppc['summary']['waste_total_usd']:,.2f}"],
        ["Zero-conversion keywords", ppc['summary']['waste_keywords_count']],
        ["High CPA keywords (above $120)", ppc['summary']['high_cpa_count']],
        ["Negative KW candidates", ppc['summary']['negative_kw_candidates_count']],
        ["Winning keywords (below $120 CPA)", ppc['summary']['winners_count']],
        ["Low CTR ads", ppc['summary']['low_ctr_ads_count']],
        ["", ""],
        ["SEO FINDINGS", ""],
        ["Quick-win keywords (pos 4-20)", seo['summary']['quick_wins_count']],
        ["Near-opportunity keywords (pos 21-50)", seo['summary']['near_opportunities_count']],
        ["Weak pages (backlinks, low visits)", seo['summary']['weak_pages_count']],
        ["Competitor gap keywords", seo['summary']['competitor_gaps_count']],
    ]
    for row in rows:
        ws_exec.append(row)
    ws_exec.column_dimensions["A"].width = 45
    ws_exec.column_dimensions["B"].width = 20
    ws_exec["A1"].font = Font(bold=True, size=14)
    ws_exec["A3"].font = Font(bold=True, color="FF" + _RED[2:])
    ws_exec["A11"].font = Font(bold=True, color="FF" + _INDIGO[2:])

    # --- PPC Waste Spend ---
    ws_waste = wb.create_sheet()
    _df_to_sheet(ws_waste, pd.DataFrame(ppc["waste_spend"]), "PPC Waste Spend", _RED)

    # --- PPC High CPA ---
    ws_hcpa = wb.create_sheet()
    _df_to_sheet(ws_hcpa, pd.DataFrame(ppc["high_cpa"]), "PPC High CPA", _ORANGE)

    # --- Negative KW Candidates ---
    ws_neg = wb.create_sheet()
    _df_to_sheet(ws_neg, pd.DataFrame(ppc["negative_kw_candidates"]), "Negative KW Candidates", _GRAY)

    # --- PPC Winners ---
    ws_win = wb.create_sheet()
    _df_to_sheet(ws_win, pd.DataFrame(ppc["winners"]), "PPC Winners", _GREEN)

    # --- SEO Quick Wins ---
    ws_qw = wb.create_sheet()
    _df_to_sheet(ws_qw, pd.DataFrame(seo["quick_wins"]), "SEO Quick Wins", _INDIGO)

    # --- SEO Near Opportunities ---
    ws_near = wb.create_sheet()
    _df_to_sheet(ws_near, pd.DataFrame(seo["near_opportunities"]), "SEO Near Opportunities", _YELLOW)

    # --- Competitor Gaps ---
    ws_gaps = wb.create_sheet()
    _df_to_sheet(ws_gaps, pd.DataFrame(seo["competitor_gaps"]), "Competitor Gaps", _INDIGO)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()
